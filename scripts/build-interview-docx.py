"""Build the interview revision handbook from its Markdown source."""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/interview-questions-and-answers.md'
OUTPUT = ROOT / 'docs/外贸CRM项目面试问答复习版.docx'
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.75)
section.bottom_margin = Inches(.7)
section.left_margin = Inches(.85)
section.right_margin = Inches(.85)
section.footer_distance = Inches(.3)


def font_style(style, size, bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    pf = style.paragraph_format
    pf.line_spacing = 1.35
    pf.space_after = Pt(7)
    pf.widow_control = True


font_style(doc.styles['Normal'], 11)
font_style(doc.styles['Title'], 23, True)
font_style(doc.styles['Subtitle'], 11)
font_style(doc.styles['Heading 1'], 15, True)
font_style(doc.styles['Heading 2'], 12, True)
font_style(doc.styles['List Bullet'], 11)
doc.styles['Subtitle'].font.italic = False
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name, before, after in [('Title', 0, 14), ('Heading 1', 17, 9), ('Heading 2', 12, 6)]:
    pf = doc.styles[name].paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = True
    pf.keep_together = True


def inline(paragraph, text):
    pattern = r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))'
    for part in re.split(pattern, text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10.5)
        elif part.startswith('[') and '](' in part:
            match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
            label, target = match.groups()
            if not target.startswith(('http:', 'https:')):
                target = (SOURCE.parent / target).resolve().as_uri()
            relationship = paragraph.part.relate_to(
                target,
                'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                is_external=True,
            )
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), relationship)
            run = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color')
            color.set(qn('w:val'), '1F4E79')
            props.append(color)
            run.append(props)
            node = OxmlElement('w:t')
            node.text = label
            run.append(node)
            link.append(run)
            paragraph._p.append(link)
        else:
            paragraph.add_run(part)


def make_table(lines):
    rows = [[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', cell) for cell in row)]
    count = len(rows[0])
    widths = [2.2, 4.6] if count == 2 else [1.6, 2.35, 2.85]
    table = doc.add_table(rows=0, cols=count)
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), 'D9D9D9')
        borders.append(border)
    props.append(borders)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        row_props = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        row_props.append(no_split)
        if row_index == 0:
            row_props.append(OxmlElement('w:tblHeader'))
        for col_index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge, size in [('top', '90'), ('bottom', '90'), ('left', '110'), ('right', '110')]:
                margin = OxmlElement(f'w:{edge}')
                margin.set(qn('w:w'), size)
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            tc_props.append(margins)
            if row_index == 0:
                fill = OxmlElement('w:shd')
                fill.set(qn('w:fill'), 'DFE8F0')
                tc_props.append(fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.25
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(10.5)
                if row_index == 0:
                    run.bold = True
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.line_spacing = Pt(3)


lines = SOURCE.read_text(encoding='utf-8').splitlines()
index = 0
while index < len(lines):
    line = lines[index].strip()
    if not line:
        index += 1
        continue
    if line.startswith('|'):
        table_lines = []
        while index < len(lines) and lines[index].strip().startswith('|'):
            table_lines.append(lines[index])
            index += 1
        make_table(table_lines)
        continue
    if line.startswith('# '):
        doc.add_paragraph('外贸 CRM 项目面试问答复习手册', 'Title')
        sub = doc.add_paragraph('62 道问答  12 个故障场景  代码复习索引', 'Subtitle')
        sub.paragraph_format.space_after = Pt(12)
    elif line.startswith('## '):
        title = re.sub(r'^[一二三四五六七八九十]+、', '', line[3:])
        title = title.replace('、', ' ').replace('，', ' ')
        doc.add_paragraph(title, 'Heading 1')
    elif line.startswith('### '):
        title = line[4:].replace('【', ' ').replace('】', ' ')
        title = re.sub(r'[？?。，、：:]', ' ', title).strip()
        question = doc.add_paragraph(title, 'Heading 2')
    elif line.startswith('>'):
        text = line[1:].strip()
        if text:
            p = doc.add_paragraph()
            inline(p, text)
    elif line.startswith('- [ ] '):
        p = doc.add_paragraph(style='List Bullet')
        inline(p, '待复习  ' + line[6:])
    else:
        p = doc.add_paragraph()
        inline(p, line)
        if re.fullmatch(r'\*\*[^*]+\*\*', line):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(3)
    index += 1

# Keep each short interview answer with its question for revision and printing.
block = []
for paragraph in doc.paragraphs:
    if paragraph.style.name in ('Heading 1', 'Heading 2'):
        if block:
            for previous in block[:-1]:
                previous.paragraph_format.keep_with_next = True
            block[-1].paragraph_format.keep_with_next = False
        block = [paragraph] if paragraph.style.name == 'Heading 2' else []
    elif block:
        block.append(paragraph)
if block:
    for previous in block[:-1]:
        previous.paragraph_format.keep_with_next = True
    block[-1].paragraph_format.keep_with_next = False

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('第 ')
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
footer.add_run(' 页').font.size = Pt(9)
doc.core_properties.title = '外贸 CRM 项目面试问答复习手册'
doc.core_properties.subject = 'Agent 开发工程师项目面试复习'
doc.core_properties.author = ''
doc.core_properties.keywords = 'CRM Agent 面试 问答 复习'
doc.save(OUTPUT)
question_count = sum(1 for p in doc.paragraphs if p.style.name == 'Heading 2' and re.match(r'Q\d+', p.text))
assert question_count == 62, question_count
print(f'Created {OUTPUT} with {question_count} questions and {len(doc.tables)} tables')
